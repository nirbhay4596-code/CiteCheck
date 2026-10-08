"""Read the text of a draft: PDF, Word (.docx) or plain text."""

import io
import unicodedata
from pathlib import Path


class UnsupportedFile(ValueError):
    pass


def extract_text(data: bytes, filename: str) -> str:
    # NFKC turns typographic ligatures (ﬁ, ﬃ) back into plain letters; PDFs are full of them
    return unicodedata.normalize("NFKC", _raw_text(data, filename))


def _raw_text(data: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        import pymupdf

        with pymupdf.open(stream=data, filetype="pdf") as pdf:
            text = "\n\n".join(page.get_text("text") for page in pdf)
        if len(text.strip()) < 50:
            raise UnsupportedFile("This PDF has no text layer (it is probably a scan). "
                                  "Run it through OCR first, or upload the Word version.")
        return text
    if suffix == ".docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        blocks = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                blocks.extend(cell.text for cell in row.cells)
        for section in document.sections:  # footnote-style citations often sit in footers
            blocks.extend(p.text for p in section.footer.paragraphs)
        return "\n\n".join(b for b in blocks if b.strip())
    if suffix in (".txt", ".md"):
        for encoding in ("utf-8", "cp1252"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
    raise UnsupportedFile(f"Can't read {suffix or 'this'} files. Upload a PDF, .docx or .txt.")


def extract_file(path: str | Path) -> str:
    path = Path(path)
    return extract_text(path.read_bytes(), path.name)
