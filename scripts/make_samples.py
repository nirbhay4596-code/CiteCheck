"""
Turn the plain-text test drafts into Word and PDF files, the way a real draft arrives.

    python scripts/make_samples.py
"""

import html
from pathlib import Path

import docx
import pymupdf
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "tests" / "drafts"
OUT = ROOT / "samples"


def is_heading(block: str) -> bool:
    return block.isupper() and len(block) < 120


def to_docx(blocks: list[str], path: Path):
    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    for block in blocks:
        p = d.add_paragraph()
        run = p.add_run(block)
        if is_heading(block):
            run.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif block.startswith("“"):
            p.paragraph_format.left_indent = Pt(36)
            p.paragraph_format.right_indent = Pt(36)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    d.save(path)


def to_pdf(blocks: list[str], path: Path):
    parts = []
    for block in blocks:
        text = html.escape(block)
        if is_heading(block):
            parts.append(f'<p style="text-align:center;font-weight:bold">{text}</p>')
        elif block.startswith("“"):
            parts.append(f'<p style="margin-left:36pt;margin-right:36pt">{text}</p>')
        else:
            parts.append(f'<p style="text-align:justify">{text}</p>')
    story = pymupdf.Story(html="".join(parts), user_css="p{font-family:serif;font-size:12pt;line-height:1.5}")
    writer = pymupdf.DocumentWriter(str(path))
    page = pymupdf.paper_rect("a4")
    body = page + (72, 72, -72, -72)
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(body)
        story.draw(device)
        writer.end_page()
    writer.close()


def main():
    OUT.mkdir(exist_ok=True)
    for src in sorted(SRC.glob("*.txt")):
        blocks = [b.strip() for b in src.read_text(encoding="utf-8").split("\n\n") if b.strip()]
        to_docx(blocks, OUT / f"{src.stem}.docx")
        to_pdf(blocks, OUT / f"{src.stem}.pdf")
        print("wrote", src.stem, ".docx and .pdf")


if __name__ == "__main__":
    main()
